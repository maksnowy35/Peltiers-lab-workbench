"""Readers for ANSYS exports and generic tables, plus column-to-role mapping.

Supported inputs
----------------
* Workbench *Table of Design Points* exported as CSV: comment lines
  ``# P1 - Current Magnitude, P2 - ...``, a header ``Name, P1, P2, ...``,
  an optional ``Units`` row and one row per design point.  When several
  analysis systems share one parameter set (e.g. variants A0/A5/A10 as
  systems A/B/C) the same output name appears several times; such
  repetitions are split into separate cases automatically.
* Workbench *Outline of All Parameters* (``ID, Parameter Name, Value, Unit``),
  one design point per file (several files are merged into rows).
* Mechanical *Tabular Data* exports (``Time [s]  Minimum [°C] ...``),
  tab-separated text or pasted from the clipboard.
* Any CSV/TSV/semicolon table or XLSX sheet with a header row; units can be
  given in brackets (``U [mV]``) or in a ``Units`` row.  Decimal commas are
  recognized for semicolon/tab separated files (European Excel locales).
"""

from __future__ import annotations

import csv
import io
import math
import re
import uuid
from dataclasses import dataclass, field

from .cases import Case, CaseMeta
from .roles import INPUT_ROLES, ROLES, suggest_role
from .units import split_name_unit, to_canonical

__all__ = [
    "ParsedColumn",
    "ParsedTable",
    "parse_bytes",
    "parse_files",
    "parse_text",
    "merge_tables",
    "suggest",
    "build_cases",
]

_PDEF = re.compile(r"\b(P\d+)\s+-\s+(.+?)(?=,\s*P\d+\s+-\s+|$)")
_PID = re.compile(r"^P\d+$")


@dataclass
class ParsedColumn:
    key: str
    name: str
    pid: str | None
    unit: str | None
    values: list[float | None]
    numeric: bool
    text: list[str] = field(default_factory=list)
    role: str = "ignore"
    case: int = 1
    confidence: float = 0.0

    def to_dict(self) -> dict:
        return {
            "key": self.key,
            "name": self.name,
            "pid": self.pid,
            "unit": self.unit,
            "values": self.values,
            "numeric": self.numeric,
            "text": self.text if not self.numeric else [],
            "role": self.role,
            "case": self.case,
            "confidence": round(self.confidence, 2),
        }

    @classmethod
    def from_dict(cls, d: dict) -> ParsedColumn:
        return cls(
            key=d["key"],
            name=d.get("name", d["key"]),
            pid=d.get("pid"),
            unit=d.get("unit"),
            values=[None if v is None else float(v) for v in d.get("values", [])],
            numeric=bool(d.get("numeric", True)),
            text=list(d.get("text", [])),
            role=d.get("role", "ignore"),
            case=int(d.get("case", 1)),
            confidence=float(d.get("confidence", 0.0)),
        )


@dataclass
class ParsedTable:
    format: str
    columns: list[ParsedColumn]
    row_labels: list[str]
    notes: list[str] = field(default_factory=list)
    sources: list[str] = field(default_factory=list)
    n_cases: int = 1
    case_names: list[str] = field(default_factory=list)
    suggested_meta: list[dict] = field(default_factory=list)

    @property
    def n_rows(self) -> int:
        return len(self.row_labels)

    def to_dict(self) -> dict:
        return {
            "format": self.format,
            "columns": [c.to_dict() for c in self.columns],
            "row_labels": self.row_labels,
            "n_rows": self.n_rows,
            "notes": self.notes,
            "sources": self.sources,
            "n_cases": self.n_cases,
            "case_names": self.case_names,
            "suggested_meta": self.suggested_meta,
        }

    @classmethod
    def from_dict(cls, d: dict) -> ParsedTable:
        cols = [ParsedColumn.from_dict(c) for c in d.get("columns", [])]
        n = max((len(c.values) for c in cols), default=0)
        return cls(
            format=d.get("format", "generic"),
            columns=cols,
            row_labels=list(d.get("row_labels") or [str(i) for i in range(n)]),
            notes=list(d.get("notes", [])),
            sources=list(d.get("sources", [])),
            n_cases=int(d.get("n_cases", 1)),
            case_names=list(d.get("case_names", [])),
        )


# ------------------------------------------------------------------ low level
def _to_float(s: object, decimal_comma: bool) -> float | None:
    if s is None:
        return None
    if isinstance(s, bool):
        return None
    if isinstance(s, (int, float)):
        return float(s) if math.isfinite(float(s)) else None
    t = str(s).strip().strip('"').strip("'").strip()
    if not t:
        return None
    if decimal_comma:
        t = t.replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        return None
    return v if math.isfinite(v) else None


def _detect_delimiter(lines: list[str]) -> str:
    best, best_score = ",", -1.0
    sample = lines[:50]
    for d in ("\t", ";", ","):
        counts = [len(next(csv.reader([ln], delimiter=d))) for ln in sample if ln.strip()]
        if not counts:
            continue
        mode = max(set(counts), key=counts.count)
        if mode < 2:
            continue
        consistency = counts.count(mode) / len(counts)
        score = consistency * 10 + min(mode, 50) / 50
        if score > best_score:
            best, best_score = d, score
    return best


def _parse_pdefs(comment_lines: list[str]) -> dict[str, str]:
    defs: dict[str, str] = {}
    for ln in comment_lines:
        body = ln.lstrip("#").strip()
        for pid, name in _PDEF.findall(body):
            defs[pid] = name.strip().rstrip(",").strip()
    return defs


def _is_numeric_row(row: list[object], decimal_comma: bool) -> float:
    cells = [c for c in row if str(c).strip()]
    if not cells:
        return 0.0
    return sum(_to_float(c, decimal_comma) is not None for c in cells) / len(cells)


def _looks_decimal_comma(rows: list[list[str]]) -> bool:
    hits = total = 0
    for r in rows[:100]:
        for c in r:
            c = c.strip()
            if re.fullmatch(r"-?\d+,\d+([eE][-+]?\d+)?", c):
                hits += 1
            if re.fullmatch(r"-?\d+(\.\d+)?([eE][-+]?\d+)?", c):
                total += 1
    return hits > 0 and hits >= total * 0.2


# ------------------------------------------------------------------ parsing
def parse_text(text: str, source: str = "pasted") -> ParsedTable:
    text = text.lstrip("﻿")
    lines = text.splitlines()
    comments = [ln for ln in lines if ln.lstrip().startswith("#")]
    body = [ln for ln in lines if ln.strip() and not ln.lstrip().startswith("#")]
    if not body:
        raise ValueError("no data rows found")
    delim = _detect_delimiter(body)
    rows = [[c.strip() for c in r] for r in csv.reader(body, delimiter=delim)]
    decimal_comma = delim in (";", "\t") and _looks_decimal_comma(rows)
    table = _parse_rows(rows, _parse_pdefs(comments), source, decimal_comma)
    if decimal_comma:
        table.notes.append("decimal commas detected")
    return table


def parse_xlsx(data: bytes, source: str) -> ParsedTable:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows: list[list[object]] = []
        comments: list[str] = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else v for v in r]
            if not any(str(c).strip() for c in cells):
                continue
            if str(cells[0]).lstrip().startswith("#"):
                comments.append(" ".join(str(c) for c in cells if str(c).strip()))
                continue
            rows.append(cells)
        if rows:
            width = max(len(r) for r in rows)
            rows = [list(r) + [""] * (width - len(r)) for r in rows]
            table = _parse_rows(rows, _parse_pdefs(comments), f"{source}:{ws.title}", False)
            return table
    raise ValueError("workbook has no data")


def parse_bytes(data: bytes, filename: str) -> ParsedTable:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data, filename)
    for enc in ("utf-8-sig", "utf-16", "cp1250", "latin-1"):
        try:
            text = data.decode(enc)
            if "\x00" in text:
                continue
            return parse_text(text, filename)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"cannot decode {filename}")


def _parse_rows(rows: list[list[object]], pdefs: dict[str, str], source: str, decimal_comma: bool) -> ParsedTable:
    width = max(len(r) for r in rows)
    rows = [list(r) + [""] * (width - len(r)) for r in rows]
    header_idx = None
    for i, r in enumerate(rows[:20]):
        if _is_numeric_row(r, decimal_comma) < 0.5 and sum(1 for c in r if str(c).strip()) >= 2:
            header_idx = i
            break
    if header_idx is None:
        header = [f"col{i + 1}" for i in range(width)]
        data = rows
    else:
        header = [str(c).strip() for c in rows[header_idx]]
        data = rows[header_idx + 1 :]
    units_row = None
    if data and str(data[0][0]).strip().lower() in ("units", "unit"):
        units_row = [str(c).strip() for c in data[0]]
        data = data[1:]

    hl = [h.lower() for h in header]
    if "parameter name" in hl and "value" in hl:
        return _pivot_outline(header, data, source, decimal_comma)

    fmt = "generic"
    if pdefs or any(_PID.match(h) for h in header):
        fmt = "workbench-design-points"
    elif any("time" in h and "[" in h for h in hl) or ("minimum" in " ".join(hl) and "maximum" in " ".join(hl)):
        fmt = "mechanical-tabular"

    columns: list[ParsedColumn] = []
    notes: list[str] = []
    label_col = None
    for j, h in enumerate(header):
        raw = [r[j] for r in data]
        vals = [_to_float(v, decimal_comma) for v in raw]
        nonempty = [v for v in raw if str(v).strip()]
        # an all-empty column (e.g. a template to be filled in) is kept as numeric so it gets a role
        numeric = (not nonempty) or sum(v is not None for v in vals) >= 0.6 * len(nonempty)
        pid = h if _PID.match(h) else None
        name = pdefs.get(h, h) if pid else h
        name, unit_hdr = split_name_unit(name or f"col{j + 1}")
        unit_cell = units_row[j] if units_row and j < len(units_row) else ""
        if unit_cell.lower() in ("units", "unit"):
            unit_cell = ""
        unit = unit_cell or unit_hdr
        if not h and not nonempty:
            continue
        col = ParsedColumn(
            key=f"c{j}",
            name=name or f"col{j + 1}",
            pid=pid,
            unit=unit,
            values=vals if numeric else [None] * len(vals),
            numeric=numeric,
            text=[str(v) for v in raw] if not numeric else [],
        )
        if not numeric and label_col is None:
            label_col = col
        columns.append(col)
    labels = label_col.text if label_col else [f"{source} #{i + 1}" for i in range(len(data))]
    if pdefs:
        notes.append(f"{len(pdefs)} parameter definitions read from the header comments")
    table = ParsedTable(format=fmt, columns=columns, row_labels=labels, notes=notes, sources=[source])
    return table


def _pivot_outline(header: list[str], data: list[list[object]], source: str, decimal_comma: bool) -> ParsedTable:
    hl = [h.lower() for h in header]
    i_id = hl.index("id") if "id" in hl else None
    i_name, i_val = hl.index("parameter name"), hl.index("value")
    i_unit = hl.index("unit") if "unit" in hl else None
    columns = []
    for k, r in enumerate(data):
        name = str(r[i_name]).strip()
        if not name:
            continue
        v = _to_float(r[i_val], decimal_comma)
        pid = str(r[i_id]).strip() if i_id is not None else None
        unit = str(r[i_unit]).strip() if i_unit is not None else None
        name, unit_hdr = split_name_unit(name)
        columns.append(
            ParsedColumn(
                key=f"c{k}",
                name=name,
                pid=pid or None,
                unit=unit or unit_hdr,
                values=[v],
                numeric=v is not None,
                text=[] if v is not None else [str(r[i_val])],
            )
        )
    return ParsedTable(
        format="workbench-parameters-outline",
        columns=columns,
        row_labels=[source],
        notes=["outline of parameters: one design point per file"],
        sources=[source],
    )


def parse_files(files: list[tuple[str, bytes]]) -> ParsedTable:
    """Parse one or more uploaded files and return a single table with suggestions."""
    return merge_tables([parse_bytes(data, name) for name, data in files])


def merge_tables(tables: list[ParsedTable]) -> ParsedTable:
    """Stack tables row-wise, matching columns by (pid or name, unit)."""
    if len(tables) == 1:
        return suggest(tables[0])
    order: list[tuple[str, str]] = []
    info: dict[tuple[str, str], ParsedColumn] = {}
    for t in tables:
        seen: dict[tuple[str, str], int] = {}
        for c in t.columns:
            base = (c.pid or c.name, c.unit or "")
            seen[base] = seen.get(base, 0) + 1
            key = (f"{base[0]}#{seen[base]}", base[1])
            if key not in info:
                info[key] = c
                order.append(key)
    columns = []
    labels: list[str] = []
    notes: list[str] = []
    for t in tables:
        labels.extend(t.row_labels)
        notes.extend(t.notes)
    for idx, key in enumerate(order):
        proto = info[key]
        vals: list[float | None] = []
        texts: list[str] = []
        for t in tables:
            seen = {}
            match = None
            for c in t.columns:
                base = (c.pid or c.name, c.unit or "")
                seen[base] = seen.get(base, 0) + 1
                if (f"{base[0]}#{seen[base]}", base[1]) == key:
                    match = c
                    break
            if match is None:
                vals.extend([None] * t.n_rows)
                texts.extend([""] * t.n_rows)
            else:
                vals.extend(match.values if match.numeric else [None] * t.n_rows)
                texts.extend(match.text if not match.numeric else [""] * t.n_rows)
        columns.append(
            ParsedColumn(
                key=f"c{idx}",
                name=proto.name,
                pid=proto.pid,
                unit=proto.unit,
                values=vals,
                numeric=proto.numeric,
                text=texts if not proto.numeric else [],
            )
        )
    fmts = {t.format for t in tables}
    merged = ParsedTable(
        format=fmts.pop() if len(fmts) == 1 else "mixed",
        columns=columns,
        row_labels=labels,
        notes=sorted(set(notes)) + [f"{len(tables)} files merged"],
        sources=[s for t in tables for s in t.sources],
    )
    return suggest(merged)


# ------------------------------------------------------------------ suggestions
def suggest(table: ParsedTable) -> ParsedTable:
    """Fill role/case suggestions and per-case metadata hints (in place)."""
    for c in table.columns:
        c.role, c.confidence = suggest_role(c.name, c.unit, c.values, c.numeric)
        if c.numeric and c.values and all(v is None for v in c.values) and table.n_rows > 1:
            # empty column (quantity not reported for this system) - keep the name but do not map it
            c.role, c.confidence = "ignore", 0.5
    outputs = [c for c in table.columns if c.numeric and c.role not in INPUT_ROLES | {"ignore", "label"}]
    counts: dict[str, list[ParsedColumn]] = {}
    for c in outputs:
        counts.setdefault(c.role, []).append(c)
    n_cases = max((len(v) for v in counts.values()), default=1)
    ranges: dict[int, list[int]] = {}
    pos = {c.key: i for i, c in enumerate(table.columns)}
    for c in table.columns:
        c.case = 0 if (c.role in INPUT_ROLES or c.role in ("label", "ignore")) else 1
    if n_cases > 1:
        for cols in counts.values():
            if len(cols) == n_cases:
                for k, c in enumerate(cols):
                    c.case = k + 1
                    ranges.setdefault(k + 1, []).append(pos[c.key])
        for cols in counts.values():
            if len(cols) == n_cases:
                continue
            for c in cols:
                p = pos[c.key]
                best = min(
                    ranges,
                    key=lambda ci: (
                        0 if min(ranges[ci]) <= p <= max(ranges[ci]) else min(abs(p - x) for x in ranges[ci])
                    ),
                )
                c.case = best
                c.confidence = min(c.confidence, 0.5)
        # inputs repeated per system (e.g. one current parameter per system)
        in_counts: dict[str, list[ParsedColumn]] = {}
        for c in table.columns:
            if c.numeric and c.role in INPUT_ROLES:
                in_counts.setdefault(c.role, []).append(c)
        for cols in in_counts.values():
            if len(cols) == n_cases:
                for k, c in enumerate(cols):
                    c.case = k + 1
        table.notes.append(f"repeated output names → {n_cases} cases (one per analysis system)")
    # pairing heuristic: 'Temperature Average' + 'T_leg_cold' -> the average is probably the hot leg ends
    for ci in range(1, n_cases + 1):
        in_case = [c for c in table.columns if c.case in (0, ci)]
        roles = {c.role for c in in_case}
        if "t_leg_cold" in roles and "t_leg_hot" not in roles:
            for c in in_case:
                if c.role == "t_avg":
                    c.role, c.confidence = "t_leg_hot", 0.4
                    table.notes.append(
                        f"case {ci}: '{c.name}' assumed to be the hot leg-end temperature – please check"
                    )
                    break
    table.n_cases = n_cases
    stem = table.sources[0].rsplit("/", 1)[-1].rsplit(".", 1)[0] if table.sources else "Imported"
    table.case_names = [stem] if n_cases == 1 else [f"{stem} – system {k + 1}" for k in range(n_cases)]
    table.suggested_meta = [_suggest_meta(table, ci) for ci in range(1, n_cases + 1)]
    return table


def _suggest_meta(table: ParsedTable, ci: int) -> dict:
    cols = {c.role: c for c in table.columns if c.case in (0, ci) and c.numeric}
    meta: dict = {"bc": "QT"}

    def spread(c: ParsedColumn) -> tuple[float, float] | None:
        vals, _ = to_canonical(c.values, c.unit, ROLES[c.role].dimension)
        v = [x for x in vals if x is not None]
        if len(v) < 2:  # a single value cannot tell a fixed from a free quantity
            return None
        return min(v), max(v)

    if "t_hot_ext" in cols:
        s = spread(cols["t_hot_ext"])
        if s and s[1] - s[0] < 0.05:
            meta.update(bc="TT", t_hot_c=round(s[0], 4))
    if "q_hot" in cols and meta["bc"] != "TT":
        s = spread(cols["q_hot"])
        if s and abs(s[1] - s[0]) <= 1e-3 * max(abs(s[1]), 1e-12):
            meta.update(bc="QT", q_hot_w=abs(s[0]))
        elif s:
            meta.update(bc="TT")
    if "t_cold_ext" in cols:
        s = spread(cols["t_cold_ext"])
        if s and s[1] - s[0] < 0.05:
            meta["t_cold_c"] = round(s[0], 4)
    return meta


# ------------------------------------------------------------------ build cases
def build_cases(
    table: ParsedTable,
    case_defs: list[dict],
    *,
    split_by: str | None = None,
    time_to_current: float | None = None,
) -> tuple[list[Case], list[str]]:
    """Turn a mapped table into cases.

    ``case_defs`` = [{"index": 1, "name": "A10-QT", "meta": {...}}, ...]; the
    column roles/case indices are read from ``table.columns``.
    """
    notes: list[str] = []
    cases: list[Case] = []
    n = table.n_rows
    split_col = next((c for c in table.columns if c.key == split_by), None) if split_by else None
    for cdef in case_defs:
        ci = int(cdef.get("index", 1))
        name = str(cdef.get("name") or f"Case {ci}")
        meta = CaseMeta(**(cdef.get("meta") or {}))
        if not meta.source:
            meta.source = ", ".join(table.sources)
        data: dict[str, list[float | None]] = {}
        for c in table.columns:
            if c.case not in (0, ci) or c.role in ("ignore", "label") or not c.numeric:
                continue
            role = ROLES.get(c.role)
            if role is None:
                continue
            conv, note = to_canonical(c.values, c.unit, role.dimension)
            if note:
                notes.append(f"{name}: {c.name}: {note}")
            if c.role in data:
                notes.append(f"{name}: second column for '{c.role}' ({c.name}) ignored")
                continue
            data[c.role] = conv
        if "current" not in data and "time" in data and time_to_current:
            data["current"] = [None if t is None else t * time_to_current for t in data["time"]]
            notes.append(f"{name}: current = time x {time_to_current} A/s")
        for q in ("q_hot", "q_cold"):
            if q in data:
                vals = [v for v in data[q] if v is not None]
                if vals and sorted(vals)[len(vals) // 2] < 0:
                    data[q] = [None if v is None else -v for v in data[q]]
                    notes.append(f"{name}: {q} sign flipped (negative reaction → positive magnitude)")
        groups: list[tuple[str, list[int]]]
        if split_col is not None:
            keys = (
                split_col.text if not split_col.numeric else ["" if v is None else f"{v:g}" for v in split_col.values]
            )
            order: dict[str, list[int]] = {}
            for i, k in enumerate(keys):
                order.setdefault(k, []).append(i)
            groups = [(f"{name} [{k}]", idx) for k, idx in order.items() if k != ""]
        else:
            groups = [(name, list(range(n)))]
        for gname, idx in groups:
            sub = {role: [vals[i] for i in idx] for role, vals in data.items()}
            keep = [
                k
                for k in range(len(idx))
                if any(sub[r][k] is not None for r in sub if r not in INPUT_ROLES)
                and ("current" not in sub or sub["current"][k] is not None)
            ]
            sub = {role: [vals[k] for k in keep] for role, vals in sub.items()}
            sort_key = "position" if meta.kind == "profile" and "position" in sub else "current"
            if sort_key in sub:
                order_idx = sorted(range(len(keep)), key=lambda k: sub[sort_key][k])
                sub = {role: [vals[k] for k in order_idx] for role, vals in sub.items()}
            if not sub or not keep:
                notes.append(f"{gname}: no usable rows – skipped")
                continue
            cases.append(Case(id=uuid.uuid4().hex[:12], name=gname, meta=meta.model_copy(), data=sub))
    return cases, notes
